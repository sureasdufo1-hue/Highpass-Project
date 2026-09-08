from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SOURCE = Path(r"C:\Users\user\Downloads\모바일분석_보고서_종합보안권고_시큐어코딩_보강본.docx")
OUTPUT = Path(r"C:\Users\user\Documents\New project\모바일분석_보고서_종합보안권고_취약점별단일표_최종본.docx")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_margins(cell, top=85, start=100, bottom=85, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def remove_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is not None:
        tr_pr.remove(cant_split)


def set_font(run, name="맑은 고딕", size=8.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)


def write_text(cell, text, *, code=False, bold=False, color="000000", style=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if style:
        p.style = style
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_font(run, "Consolas" if code else "맑은 고딕", 7.3 if code else 8.5, bold, color)
    return p


def next_sibling(el):
    return el.getnext()


def paragraph_text(el):
    if el.tag != qn("w:p"):
        return ""
    from docx.text.paragraph import Paragraph
    return Paragraph(el, el.getparent()).text


def table_from_el(el):
    from docx.table import Table
    return Table(el, doc)


doc = Document(SOURCE)
body = doc._element.body
headings = [p for p in doc.paragraphs if p.style.name == "Heading 2" and re.match(r"^6\.(?:[1-9]|1[0-3])\s+MOB-", p.text.strip())]
if len(headings) != 13:
    raise RuntimeError(f"Expected 13 chapter-6 headings, found {len(headings)}")

for index, heading in enumerate(headings, 1):
    start = heading._p
    end = headings[index]._p if index < len(headings) else next(
        p._p for p in doc.paragraphs if p.style.name == "Heading 2" and p.text.strip().startswith("6.14 종합 개선 실행계획")
    )
    start_idx = body.index(start)
    end_idx = body.index(end)
    nodes = list(body)[start_idx:end_idx]

    old_table_el = next((n for n in nodes if n.tag == qn("w:tbl")), None)
    if old_table_el is None:
        raise RuntimeError(f"No recommendation table for {heading.text}")
    old_table = table_from_el(old_table_el)
    base = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in old_table.rows if len(row.cells) >= 2}

    paras = [paragraph_text(n).strip() for n in nodes if n.tag == qn("w:p")]
    code_i = next(i for i, t in enumerate(paras) if t.startswith("시큐어 코딩 / 설정 예시"))
    cmd_i = next(i for i, t in enumerate(paras) if t == "재점검 명령어")
    code_heading = paras[code_i]
    code_type = code_heading.partition("(")[2].rstrip(")") or "Code"
    code_text = paras[code_i + 1]
    code_explain = paras[code_i + 2]
    commands = paras[cmd_i + 1]
    expected = paras[cmd_i + 2]
    if expected.startswith("기대 결과:"):
        expected = expected[len("기대 결과:"):].strip()

    new_table = doc.add_table(rows=0, cols=2)
    new_table.style = old_table.style
    new_table.autofit = False
    # 원래 표의 테두리·폭·정렬 속성을 그대로 복제한다.
    new_table._tbl.remove(new_table._tbl.tblPr)
    new_table._tbl.insert(0, deepcopy(old_table._tbl.tblPr))
    if old_table._tbl.tblGrid is not None:
        existing_grid = new_table._tbl.tblGrid
        if existing_grid is not None:
            new_table._tbl.remove(existing_grid)
        new_table._tbl.insert(1, deepcopy(old_table._tbl.tblGrid))

    title_row = new_table.add_row()
    title_cell = title_row.cells[0].merge(title_row.cells[1])
    shade(title_cell, "404040")
    set_margins(title_cell, 110, 120, 110, 120)
    title_p = write_text(title_cell, heading.text.strip(), bold=True, color="FFFFFF", style="Heading 2")
    title_p.paragraph_format.keep_with_next = True
    set_repeat_header(title_row)

    ordered = [
        ("진단 결과", base["진단 결과"], False),
        ("기준 위험도", base["기준 위험도"], False),
        ("현재 문제", base["현재 문제"], False),
        ("보안상 영향", base["보안상 영향"], False),
        ("개선 목표", base["개선 목표"], False),
        ("적용 방법", base["적용 방법"], False),
        ("시큐어 코딩 / 설정 예시", f"[{code_type}]\n{code_text}", True),
        ("코드 설명", code_explain, False),
        ("재점검 명령어", commands, True),
        ("기대 결과", expected, False),
        ("완료 기준", base["완료 기준"], False),
        ("우선순위", base["우선순위"], False),
    ]
    for label, value, is_code in ordered:
        row = new_table.add_row()
        remove_no_split(row)
        left, right = row.cells
        set_cell_width(left, 1720)
        set_cell_width(right, 7000)
        set_margins(left)
        set_margins(right)
        shade(left, "D9D9D9")
        shade(right, "FFFFFF")
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        write_text(left, label, bold=True)
        write_text(right, value, code=is_code)

    # 표 시작 전의 빈 문단에 페이지 나누기를 넣고, 새 표를 기존 위치로 이동한다.
    page_break = doc.add_paragraph()
    if index > 1:
        page_break.add_run().add_break(WD_BREAK.PAGE)
    start.addprevious(page_break._p)
    start.addprevious(new_table._tbl)
    for node in nodes:
        if node.getparent() is body:
            body.remove(node)

# 목차 캐시는 Word 자동 갱신 전에도 6.1~6.13을 유지한다.
toc_texts = [p.text for p in doc.paragraphs if p.style.name.lower().startswith("toc")]
for heading in headings:
    number = heading.text.split()[0]
    if not any(t.startswith(number + " ") for t in toc_texts):
        raise RuntimeError(f"TOC entry missing before save: {heading.text}")

doc.core_properties.title = "모바일 애플리케이션 보안 취약점 분석 보고서 - 취약점별 단일표 최종본"
doc.save(OUTPUT)
print(OUTPUT)
